"""Document ingestion.

Converts an uploaded complaint artefact into plain text the agent can reason
over. Production-grade OCR is explicitly out of scope, so scanned images are
detected and reported honestly rather than silently returning empty text.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field

from app.core.config import settings

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".eml", ".msg", ".html", ".htm"}

_WS = re.compile(r"[ \t ]+")
_BLANKS = re.compile(r"\n{3,}")
_TAGS = re.compile(r"<[^>]+>")


@dataclass
class ParsedDocument:
    text: str
    filename: str
    kind: str
    pages: int = 1
    warnings: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


# Field labels that appear on a line of their own in a tabular document.
# PDF text extraction emits one line per visual cell, so "Batch No." and
# "LEV2604A" arrive as two separate lines and every label-based rule misses.
# Re-pairing them here means the rest of the pipeline sees one canonical shape
# regardless of whether the source was an e-mail, a Word form or a PDF table.
LABEL_STEMS: frozenset[str] = frozenset({
    "product", "product name", "product concerned", "material", "material name",
    "strength", "product strength", "strength / grade", "grade", "dosage form",
    "dosage", "presentation", "pack", "pack size",
    "batch", "batch no", "batch no.", "batch number", "lot", "lot no", "lot no.",
    "lot number", "b.no", "b no",
    "mfg date", "mfg. date", "manufacturing date", "date of manufacture", "mfd",
    "exp date", "exp. date", "expiry", "expiry date", "expiration date",
    "retest date", "use before",
    "quantity", "quantity affected", "quantity received", "qty", "units affected",
    "sample available", "sample availability", "sample",
    "complaint raised by", "reported by", "raised by", "complainant",
    "contact person", "contact", "customer", "customer name", "organisation",
    "organization", "country", "market", "destination",
    "reference", "ref", "ref no", "complaint no", "complaint number",
    "date of complaint", "complaint date", "date received", "date",
    "nature of complaint", "complaint type", "defect", "defect type",
    "severity", "classification", "priority",
})

_TRAILING = " \t:.-–—"


def pair_orphan_labels(text: str) -> str:
    """Join a bare label line to the value line beneath it."""
    lines = text.split("\n")
    out: list[str] = []
    index = 0
    while index < len(lines):
        current = lines[index]
        stem = current.strip().strip(_TRAILING).lower()
        nxt = lines[index + 1].strip() if index + 1 < len(lines) else ""
        if (
            stem in LABEL_STEMS
            and ":" not in current
            and nxt
            and len(nxt) <= 110
            and nxt.strip().strip(_TRAILING).lower() not in LABEL_STEMS
        ):
            out.append(f"{current.strip().strip(_TRAILING)}: {nxt}")
            index += 2
            continue
        out.append(current)
        index += 1
    return "\n".join(out)


def normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS.sub(" ", text)
    text = _BLANKS.sub("\n\n", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = pair_orphan_labels(text)
    return text.strip()[: settings.max_document_chars]


def _parse_pdf(data: bytes) -> tuple[str, int, list[str]]:
    from pypdf import PdfReader

    warnings: list[str] = []
    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"Page skipped: {exc}")
    text = "\n\n".join(chunks)
    if len(text.strip()) < 40:
        warnings.append(
            "This PDF contains little or no extractable text — it is most likely a "
            "scan. OCR is out of scope; paste the complaint text instead."
        )
    return text, len(reader.pages), warnings


def _parse_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_eml(data: bytes) -> tuple[str, dict]:
    import email
    from email import policy

    message = email.message_from_bytes(data, policy=policy.default)
    meta = {
        "from": str(message.get("From") or ""),
        "to": str(message.get("To") or ""),
        "subject": str(message.get("Subject") or ""),
        "date": str(message.get("Date") or ""),
    }
    body = ""
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                body += part.get_content()
            elif part.get_content_type() == "text/html" and not body:
                body += _TAGS.sub(" ", part.get_content())
    else:
        body = message.get_content()

    header = "\n".join(f"{k.title()}: {v}" for k, v in meta.items() if v)
    return f"{header}\n\n{body}", meta


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    if len(data) > settings.max_upload_bytes:
        raise ValueError(
            f"File exceeds the {settings.max_upload_bytes // (1024 * 1024)} MB limit."
        )

    lower = filename.lower()
    suffix = lower[lower.rfind(".") :] if "." in lower else ""
    warnings: list[str] = []
    pages = 1
    metadata: dict = {}

    if suffix == ".pdf":
        raw, pages, warnings = _parse_pdf(data)
        kind = "pdf"
    elif suffix == ".docx":
        raw, kind = _parse_docx(data), "docx"
    elif suffix in {".eml", ".msg"}:
        raw, metadata = _parse_eml(data)
        kind = "email"
    elif suffix in {".html", ".htm"}:
        raw, kind = _TAGS.sub(" ", data.decode("utf-8", "replace")), "html"
    elif suffix in {".txt", ".md", ""}:
        raw, kind = data.decode("utf-8", "replace"), "text"
    else:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: "
            + ", ".join(sorted(SUPPORTED_EXTENSIONS))
        )

    text = normalise(raw)
    if not text:
        warnings.append("No readable text could be extracted from this file.")

    return ParsedDocument(
        text=text,
        filename=filename,
        kind=kind,
        pages=pages,
        warnings=warnings,
        metadata=metadata,
    )
